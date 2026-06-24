from __future__ import annotations

import re
import sys
from dataclasses import dataclass
from pathlib import Path

# Allow importing from the project root (ecb_pdf_to_json, pdf_extract)
_ROOT = Path(__file__).parent.parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))


@dataclass
class Chunk:
    chunk_id: str
    text: str
    char_start: int
    char_end: int
    page_hint: int | None = None


DOMAIN_VOCAB: dict[str, list[str]] = {
    "credit_risk": [
        "irb", "lgd", "pd", "ead", "probability", "default", "rating",
        "grade", "calibration", "downturn", "mortgage", "conversion",
        "factor", "retail", "corporate", "exposure", "slotting",
    ],
    "market_risk_crr2": [
        "var", "stressed", "irc", "rnime", "backtesting", "vega",
        "gamma", "trading", "incremental", "risk", "charge",
    ],
    "market_risk_crr3": [
        "frtb", "drc", "ssrm", "expected", "shortfall", "sensitivities",
        "nmrf", "rrao", "curvature", "delta", "bucket",
    ],
    "ccr": [
        "epe", "imm", "netting", "margin", "collateral", "alpha",
        "cva", "counterparty", "mpor", "eepe", "mtm",
    ],
}

_SENTENCE_END = re.compile(r"[.!?]\s+[A-Z]")


def load_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".pdf":
        from ecb_pdf_to_json import extract_blocks, segment_paragraphs
        blocks = extract_blocks(path)
        paragraphs = segment_paragraphs(blocks)
        return "\n\n".join(p.text for p in paragraphs)
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        parts: list[str] = []
        for ws in wb.worksheets:
            parts.append(f"=== Sheet: {ws.title} ===")
            for row in ws.iter_rows(values_only=True):
                line = "  ".join(str(c) for c in row if c is not None)
                if line.strip():
                    parts.append(line)
        wb.close()
        return "\n".join(parts)
    raise ValueError(f"Unsupported file type: {suffix!r}. Use .pdf, .txt, or .xlsx")


def chunk_text(
    text: str,
    chunk_size: int = 2000,
    overlap: int = 200,
) -> list[Chunk]:
    chunks: list[Chunk] = []
    start = 0
    idx = 0
    n = len(text)

    while start < n:
        end = min(start + chunk_size, n)
        if end < n:
            # Try to break at a sentence boundary within the last 20% of the chunk
            search_from = start + int(chunk_size * 0.8)
            match = None
            for m in _SENTENCE_END.finditer(text, search_from, end):
                match = m
            if match:
                end = match.start() + 1  # include the period

        chunk_text_str = text[start:end].strip()
        if chunk_text_str:
            chunks.append(Chunk(
                chunk_id=f"chunk_{idx:04d}",
                text=chunk_text_str,
                char_start=start,
                char_end=end,
            ))
            idx += 1

        next_start = end - overlap
        if next_start <= start:
            next_start = end
        start = next_start

    return chunks


def detect_scope(text: str) -> list[str]:
    from backend.agent.bm25_index import BM25Index

    sample = text[:5000]
    idx = BM25Index()
    domain_names = list(DOMAIN_VOCAB.keys())
    idx.add_documents([" ".join(DOMAIN_VOCAB[d]) for d in domain_names])

    results = idx.query(sample, top_k=len(domain_names))
    if not results:
        return list(DOMAIN_VOCAB.keys())

    top_score = results[0][1]
    threshold = top_score * 0.20
    detected = [domain_names[doc_idx] for doc_idx, score in results if score >= threshold]
    return detected if detected else list(DOMAIN_VOCAB.keys())
