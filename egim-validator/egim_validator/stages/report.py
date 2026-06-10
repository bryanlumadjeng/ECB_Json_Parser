"""Stage 6 — reporting: merge agent verdicts with validator overrides and export.

Exports: Excel traceability matrix (Matrix + Summary sheets), JSON, Markdown summary.
Validator overrides always win over agent verdicts in the FINAL columns; the agent's
original verdict is preserved alongside for the audit trail.
"""

from __future__ import annotations

import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..schemas import AssessmentRecord, CoverageStatus, Override, ReviewState, final_status

STATUS_FILL = {
    CoverageStatus.covered: "C6EFCE",
    CoverageStatus.partially_covered: "FFEB9C",
    CoverageStatus.gap: "FFC7CE",
    CoverageStatus.unclear: "D9D2E9",
    CoverageStatus.not_applicable: "EEEEEE",
}

COLUMNS = [
    ("Ref", 12),
    ("Chapter", 24),
    ("Section", 28),
    ("Para", 7),
    ("Kind", 12),
    ("Guide paragraph", 60),
    ("Applicable", 11),
    ("Applicability reason", 40),
    ("Agent status", 17),
    ("Final status", 17),
    ("Review state", 13),
    ("Confidence", 11),
    ("Evidence (quotes)", 70),
    ("Evidence sections", 18),
    ("Quotes verified", 14),
    ("Rationale", 70),
    ("Open questions", 50),
    ("Flags", 40),
    ("Validator notes", 50),
]


def build_rows(
    records: list[AssessmentRecord], overrides: dict[str, Override]
) -> list[dict[str, object]]:
    rows = []
    for r in sorted(records, key=lambda x: x.ref):
        ov = overrides.get(r.ref)
        fs = final_status(r, ov)
        rows.append(
            {
                "Ref": r.ref,
                "Chapter": r.chapter or "",
                "Section": r.section or "",
                "Para": r.paragraph_number or "",
                "Kind": r.kind.value,
                "Guide paragraph": r.paragraph_text,
                "Applicable": "yes" if r.applicability.applicable else "no",
                "Applicability reason": r.applicability.reason,
                "Agent status": r.status.value if r.status else "",
                "Final status": fs.value if fs else "",
                "Review state": (ov.review_state.value if ov else ReviewState.pending.value),
                "Confidence": r.confidence.value if r.confidence else "",
                "Evidence (quotes)": "\n".join(f'[{e.section_id}] "{e.quote}"' for e in r.evidence),
                "Evidence sections": ", ".join(sorted({e.section_id for e in r.evidence})),
                "Quotes verified": {True: "yes", False: "NO", None: ""}[r.all_evidence_verified],
                "Rationale": r.rationale,
                "Open questions": "\n".join(r.open_questions),
                "Flags": "; ".join(r.flags),
                "Validator notes": ov.notes if ov else "",
            }
        )
    return rows


def to_excel(rows: list[dict[str, object]], path: str | Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Matrix"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3864")
    for col, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    wrap = Alignment(wrap_text=True, vertical="top")
    final_col = next(i for i, (n, _) in enumerate(COLUMNS, start=1) if n == "Final status")
    for row_idx, row in enumerate(rows, start=2):
        for col, (name, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col, value=row.get(name, ""))
            cell.alignment = wrap
        try:
            status = CoverageStatus(rows[row_idx - 2]["Final status"])
            ws.cell(row=row_idx, column=final_col).fill = PatternFill(
                "solid", fgColor=STATUS_FILL[status]
            )
        except ValueError:
            pass

    summary = wb.create_sheet("Summary")
    counts = Counter((row["Chapter"], row["Final status"] or "(unassessed)") for row in rows)
    chapters = sorted({c for c, _ in counts})
    statuses = [s.value for s in CoverageStatus] + ["(unassessed)"]
    summary.cell(row=1, column=1, value="Chapter").font = Font(bold=True)
    for j, s in enumerate(statuses, start=2):
        summary.cell(row=1, column=j, value=s).font = Font(bold=True)
        summary.column_dimensions[get_column_letter(j)].width = 18
    summary.column_dimensions["A"].width = 40
    for i, chapter in enumerate(chapters, start=2):
        summary.cell(row=i, column=1, value=chapter)
        for j, s in enumerate(statuses, start=2):
            summary.cell(row=i, column=j, value=counts.get((chapter, s), 0))

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def to_json(rows: list[dict[str, object]], path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(rows, ensure_ascii=False, indent=1), encoding="utf-8")
    return path


def to_markdown_summary(
    rows: list[dict[str, object]], path: str | Path, model_name: str = ""
) -> Path:
    applicable = [r for r in rows if r["Applicable"] == "yes"]
    counts = Counter(r["Final status"] for r in applicable)
    gaps = [r for r in applicable if r["Final status"] == CoverageStatus.gap.value]
    flagged = [r for r in rows if r["Flags"]]

    lines = [
        f"# Validation summary — {model_name}".rstrip(" —"),
        "",
        f"Generated: {datetime.now(timezone.utc).isoformat(timespec='seconds')}",
        "",
        f"Applicable EGIM paragraphs: **{len(applicable)}** of {len(rows)}",
        "",
        "| Final status | Count |",
        "|---|---|",
    ]
    for status in CoverageStatus:
        lines.append(f"| {status.value} | {counts.get(status.value, 0)} |")
    lines += ["", "## Gaps", ""]
    if gaps:
        for r in gaps:
            lines.append(f"- **{r['Ref']}** ({r['Chapter']} / {r['Section']}): {str(r['Guide paragraph'])[:200]}")
    else:
        lines.append("None.")
    lines += ["", "## Flagged for review", ""]
    if flagged:
        for r in flagged:
            lines.append(f"- **{r['Ref']}**: {r['Flags']}")
    else:
        lines.append("None.")

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path
